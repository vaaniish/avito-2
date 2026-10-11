from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "artifacts" / "discovery"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_PATH = OUT_DIR / "voprosy-dlya-obsuzhdeniya-kontseptsii.docx"

FONT = "Calibri"
INK = "0B2545"
BLUE = "2E74B5"
BLUE_DARK = "1F4D78"
MUTED = "5C6773"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
LIGHT_GREEN = "EAF4EE"
GREEN = "285943"
LIGHT_GOLD = "FBF4DE"
GOLD = "7A5A00"
WHITE = "FFFFFF"
BORDER = "CCD6E0"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=TABLE_INDENT_DXA):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run(run, size=11, color=INK, bold=False, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def set_paragraph(p, before=0, after=6, line=1.25, keep=False):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep


def add_para(doc, text="", size=11, color=INK, bold=False, italic=False,
             before=0, after=6, line=1.25, align=None, keep=False):
    p = doc.add_paragraph()
    set_paragraph(p, before, after, line, keep)
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_rich_para(doc, parts, before=0, after=6, line=1.25, keep=False):
    p = doc.add_paragraph()
    set_paragraph(p, before, after, line, keep)
    for text, kwargs in parts:
        run = p.add_run(text)
        set_run(run, **kwargs)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    p.paragraph_format.keep_with_next = True
    return p


def add_bullet(doc, text, level=0, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_paragraph(p, after=4, line=1.25)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run(r2)
    else:
        set_run(p.add_run(text))
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_paragraph(p, after=4, line=1.25)
    set_run(p.add_run(text))
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE, label_color=BLUE_DARK):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    set_table_borders(table, color=fill, size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    set_paragraph(p, after=0, line=1.2)
    r1 = p.add_run(f"{label} ")
    set_run(r1, size=10.5, color=label_color, bold=True)
    r2 = p.add_run(text)
    set_run(r2, size=10.5, color=INK)
    add_para(doc, "", after=2)
    return table


def add_question(doc, number, title, prompt, options, recommendation, why=None, current=None):
    add_heading(doc, f"{number}. {title}", 2)
    if why:
        add_rich_para(
            doc,
            [("Почему это важно. ", {"size": 10.5, "color": BLUE_DARK, "bold": True}),
             (why, {"size": 10.5, "color": MUTED})],
            after=5,
            line=1.2,
        )
    add_para(doc, prompt, keep=True)
    for label, text in options:
        add_bullet(doc, f"{label} — {text}", bold_prefix=f"{label} —")
    if current:
        add_callout(doc, "Сейчас в проекте:", current, fill=LIGHT_GRAY, label_color=MUTED)
    add_callout(doc, "Предлагаемый вариант:", recommendation, fill=LIGHT_GREEN, label_color=GREEN)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        1: (16, BLUE, 18, 10),
        2: (13, BLUE, 14, 7),
        3: (12, BLUE_DARK, 10, 5),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Страница ")
    set_run(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])


def add_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph(p, after=0, line=1.0)
    r1 = p.add_run("DISCOVERY BRIEF")
    set_run(r1, size=8.5, color=BLUE, bold=True)
    r2 = p.add_run("   |   Площадка по технике и услугам")
    set_run(r2, size=8.5, color=MUTED)

    footer = section.footer
    p = footer.paragraphs[0]
    set_paragraph(p, after=0, line=1.0)
    add_page_field(p)


def add_metric_strip(doc):
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [3120, 3120, 3120])
    set_table_borders(table, color=BORDER, size="4")
    values = [
        ("ФОРМАТ", "Совместное обсуждение"),
        ("ВРЕМЯ", "Около 60 минут"),
        ("РЕЗУЛЬТАТ", "Зафиксированные решения"),
    ]
    for idx, (label, value) in enumerate(values):
        cell = table.cell(0, idx)
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph(p, after=1, line=1.0)
        r1 = p.add_run(label + "\n")
        set_run(r1, size=8, color=BLUE, bold=True)
        r2 = p.add_run(value)
        set_run(r2, size=9.5, color=INK, bold=True)
    add_para(doc, "", after=1)


def add_summary_table(doc):
    add_heading(doc, "Итог встречи", 1)
    add_para(doc, "После разговора достаточно зафиксировать выбранный вариант по каждому пункту. Детальный интерфейс и техническая архитектура проектируются отдельно.")
    rows = [
        "Главная проблема и ценность площадки",
        "Роль площадки в сделке",
        "Главное направление запуска",
        "Роли и возможности одного аккаунта",
        "Способ входа и подтверждения контакта",
        "Допустимые типы мастеров",
        "Проверка и платная активация мастера",
        "Типы публикаций",
        "Механизм запроса и отклика",
        "Чат, гости и передача контактов",
        "Обещания и ответственность площадки",
        "Модерация, жалобы и санкции",
        "Профиль и кабинет мастера",
        "Границы первой версии",
    ]
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [4320, 5040])
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, text in enumerate(("Решение", "Что согласовано / что уточнить")):
        cell = hdr.cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        set_paragraph(p, after=0, line=1.1)
        set_run(p.add_run(text), size=9.5, color=BLUE_DARK, bold=True)
    for item in rows:
        cells = table.add_row().cells
        set_cell_width(cells[0], 4320)
        set_cell_width(cells[1], 5040)
        for cell in cells:
            set_cell_margins(cell, top=90, bottom=90)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p0 = cells[0].paragraphs[0]
        set_paragraph(p0, after=0, line=1.1)
        set_run(p0.add_run(item), size=9.5, color=INK)
        p1 = cells[1].paragraphs[0]
        set_paragraph(p1, after=0, line=1.1)
        set_run(p1.add_run(""), size=9.5, color=MUTED)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)
configure_styles(doc)
add_header_footer(section)

# First-page workshop header. Named overrides: title 28 pt, subtitle 13 pt.
add_para(doc, "АНАЛИТИЧЕСКАЯ ВСТРЕЧА", size=9, color=BLUE, bold=True, after=8)
add_para(doc, "Вопросы для обсуждения\nконцепции проекта", size=28, color=INK, bold=True, after=7, line=1.0)
add_para(doc, "Специализированная доска объявлений по технике, мастерам и услугам", size=13, color=MUTED, after=18, line=1.15)
add_metric_strip(doc)
add_callout(
    doc,
    "Цель:",
    "согласовать фундаментальные правила площадки. Расположение кнопок, внешний вид форм, технологии и внутренняя архитектура на этой встрече не обсуждаются.",
    fill=LIGHT_GOLD,
    label_color=GOLD,
)

add_heading(doc, "Исходное понимание", 1)
add_para(doc, "Проект рассматривается как специализированная доска объявлений, а не универсальный магазин. Предлагаются три направления:")
add_bullet(doc, "продажа новой, б/у и восстановленной техники;")
add_bullet(doc, "предложения услуг мастеров и сервисных центров;")
add_bullet(doc, "запросы клиентов на ремонт, диагностику и обслуживание.")
add_para(doc, "Заказчик выбирает бизнес-модель и правила взаимодействия. Аналитик предлагает варианты, фиксирует решения и затем самостоятельно проектирует интерфейс и техническую реализацию.", color=MUTED, italic=True, after=2)

add_heading(doc, "I. Основа продукта", 1)

add_question(
    doc, 1, "Главная проблема",
    "Какую проблему площадка должна решать заметно лучше существующих досок объявлений?",
    [
        ("A", "помогать быстро находить подходящего мастера"),
        ("B", "создавать доверие к б/у и восстановленной технике"),
        ("C", "объединять ремонт, восстановление и продажу техники"),
    ],
    "Вариант C — единая идея «вторая жизнь техники»: устройство можно купить, продать, восстановить или отдать в ремонт.",
    why="Ответ определяет позиционирование, главную страницу и приоритеты запуска.",
)

add_question(
    doc, 2, "Роль площадки в сделке",
    "Где заканчивается участие сайта?",
    [
        ("A", "сайт публикует объявления и передаёт контакты"),
        ("B", "сайт предоставляет объявления, отклики и чат, но не принимает оплату"),
        ("C", "товары оплачиваются через сайт, услуги — напрямую"),
        ("D", "все расчёты и споры проходят через платформу"),
    ],
    "Для первой версии — вариант B. Безопасную сделку и приём оплаты услуг добавлять после проверки спроса.",
    why="От этого зависят платежи, возвраты, арбитраж, юридическая ответственность и объём разработки.",
)

add_question(
    doc, 3, "Главное направление запуска",
    "На чём делаем первый акцент?",
    [
        ("A", "услуги и поиск мастеров"),
        ("B", "продажа б/у и восстановленной техники"),
        ("C", "оба направления одинаково важны"),
    ],
    "Сохранить оба направления, но выбрать одно главным для первого экрана и продвижения. Предварительно — услуги и запросы как отличие от обычного магазина.",
    why="Запустить всё с одинаковой глубиной силами одного разработчика рискованно.",
)

add_heading(doc, "II. Пользователи и доступ", 1)

add_question(
    doc, 4, "Роли в одном аккаунте",
    "Может ли один человек одновременно покупать, продавать свою технику и оказывать услуги?",
    [
        ("A", "для каждой роли нужен отдельный аккаунт"),
        ("B", "один аккаунт, к которому подключаются возможности продавца и мастера"),
        ("C", "обычные пользователи и полностью отдельные профессиональные аккаунты"),
    ],
    "Вариант B. Мастер сам может быть покупателем и продавцом; повторная регистрация ему не нужна.",
    why="Это определяет модель пользователей, права доступа и структуру личного кабинета.",
)

add_question(
    doc, 5, "Регистрация и вход",
    "Каким способом обычный пользователь создаёт аккаунт и возвращается на сайт?",
    [
        ("A", "email и пароль с подтверждением почты"),
        ("B", "телефон и одноразовый код без пароля"),
        ("C", "VK ID или другой социальный вход"),
        ("D", "несколько способов на выбор"),
    ],
    "Для MVP — подтверждённые email или телефон; перед публикацией и перепиской контакт обязателен. VK ID оставить дополнительным способом после запуска.",
    why="Авторизация должна быть простой для клиента, но позволять восстановить доступ и ограничивать злоупотребления.",
    current="регистрация по имени, username, email и паролю; вход по email и паролю. Подтверждения контакта и социальных входов нет.",
)

add_question(
    doc, 6, "Кто может стать мастером",
    "Какие статусы исполнителей допускаются?",
    [
        ("A", "только физические лица"),
        ("B", "только самозанятые и ИП"),
        ("C", "только ИП и организации"),
        ("D", "физлица, самозанятые, ИП и организации с публичным отображением статуса"),
    ],
    "Вариант D. Разрешить разные статусы, но использовать разные уровни проверки и явно показывать их клиенту.",
    why="Тип исполнителя влияет на доверие, документы, платежи и возможные гарантии.",
    current="партнёрский онбординг рассчитан только на ИП и юридические лица с проверкой ИНН и регистрационных данных.",
)

add_question(
    doc, 7, "Проверка и платная активация мастера",
    "За какое действие или преимущество мастер должен платить?",
    [
        ("A", "за само создание аккаунта"),
        ("B", "за проверку и активацию профессионального профиля"),
        ("C", "за подписку и доступ к запросам"),
        ("D", "за отдельные отклики или контакты"),
        ("E", "только за продвижение"),
    ],
    "Не брать плату за пустой аккаунт. Первым мастерам дать тестовый период, затем монетизировать профессиональную активацию, доступ к заявкам или продвижение.",
    why="Нужно определить не стоимость, а момент оплаты и возможности, которые она открывает.",
)

add_heading(doc, "III. Публикации и взаимодействие", 1)

add_question(
    doc, 8, "Типы публикаций",
    "Достаточно ли трёх самостоятельных типов?",
    [
        ("A", "«Продам технику»"),
        ("B", "«Окажу услугу»"),
        ("C", "«Ищу мастера»"),
    ],
    "Сохранить все три типа и не смешивать их в одну форму: у них разные данные и жизненный цикл.",
    why="Это фундаментальная структура каталога, поиска и базы данных.",
)

add_question(
    doc, 9, "Запрос клиента и отклики",
    "Как клиент должен получать предложения мастеров?",
    [
        ("A", "сам выбирает мастера из каталога"),
        ("B", "публикует запрос, на который откликаются мастера"),
        ("C", "используются оба способа"),
        ("D", "площадка сама назначает исполнителя"),
    ],
    "Вариант C. Запрос видят подходящие проверенные мастера; в отклике указываются цена, срок, гарантия и комментарий; клиент выбирает одного исполнителя.",
    why="Это центральный бизнес-процесс сервисной части проекта.",
)

add_question(
    doc, 10, "Чат, гости и контакты",
    "Когда стороны могут начать общение и получить контакты?",
    [
        ("A", "чата нет, контакты видны сразу"),
        ("B", "чат есть, контакты можно передавать сразу"),
        ("C", "до выбора мастера — чат, после выбора можно открыть контакты"),
        ("D", "контакты всегда скрыты"),
    ],
    "Вариант C. Просмотр доступен всем; гостю для запроса или сообщения выдаётся сессия после подтверждения телефона/email. Полностью анонимные сообщения не разрешаются.",
    why="Решение влияет на безопасность, спам, ценность платного доступа мастера и способность восстановить диалог.",
)

add_heading(doc, "IV. Доверие и контроль", 1)

add_question(
    doc, 11, "Что площадка обещает пользователю",
    "Какой уровень доверия обеспечивает сервис?",
    [
        ("A", "только размещение объявлений"),
        ("B", "подтверждение личности или статуса мастера"),
        ("C", "проверка квалификации и восстановленной техники"),
        ("D", "гарантия результата и разрешение коммерческих споров"),
    ],
    "Для старта — вариант B с прозрачными отметками проверки. Гарантию качества обещать только после появления регламентов и операционной возможности её обеспечить.",
    why="Именно обещания площадки определяют её ответственность, а не само наличие кнопки «Пожаловаться».",
)

add_question(
    doc, 12, "Модерация, жалобы и санкции",
    "Как контролировать пользовательские публикации?",
    [
        ("A", "всё проверяется вручную до публикации"),
        ("B", "работают только автоматические правила"),
        ("C", "правила оценивают риск, пользователи жалуются, спорные случаи решает администратор"),
        ("D", "решение полностью принимает модель искусственного интеллекта"),
    ],
    "Вариант C. Начать со структурированных полей, правил, антиспама, жалоб и ручного решения серьёзных случаев. ИИ позднее использовать как помощника модератора, а не как единственный источник блокировки.",
    why="Нужно одновременно быстро публиковать нормальные объявления и не применять тяжёлые санкции по ложным срабатываниям.",
    current="уже обнаруживаются контакты, обход площадки, мат, спам, подозрительные цены и часть проблем изображений. Есть жалобы и лестница санкций: предупреждение, 3 дня, 30 дней, постоянная блокировка.",
)

add_para(doc, "Для отдельного согласования", size=10.5, color=BLUE_DARK, bold=True, before=2, after=4, keep=True)
add_bullet(doc, "на что можно жаловаться: объявление, сообщение, профиль, отклик, отзыв;")
add_bullet(doc, "какие нарушения требуют немедленного скрытия;")
add_bullet(doc, "кто принимает решение и возможна ли апелляция;")
add_bullet(doc, "какие материалы и способы передачи контактов запрещены.")

add_heading(doc, "V. Кабинет и границы запуска", 1)

add_question(
    doc, 13, "Профиль и кабинет мастера",
    "Достаточно ли существующих возможностей партнёра или профессиональному профилю нужны отдельные разделы?",
    [
        ("A", "сохранить кабинет практически без изменений"),
        ("B", "добавить к товарному кабинету функции мастера"),
        ("C", "создать полностью отдельные кабинеты продавца и мастера"),
    ],
    "Вариант B. Один аккаунт и общий профиль, но разделы показываются по подключённым возможностям: товары, услуги, запросы, отклики, чаты, работы, отзывы, документы и профессиональный доступ.",
    why="Профиль должен отражать реальную деятельность, не заставляя мастера заводить несколько аккаунтов.",
    current="публично показываются имя, город, проверка, время ответа, рейтинг, отзывы и объявления. В кабинете есть профиль, адреса, покупки, избранное, объявления, финансы, вопросы и заказы.",
)

add_question(
    doc, 14, "Границы первой версии",
    "Что ограничиваем на первом запуске?",
    [
        ("A", "несколько категорий техники и один регион"),
        ("B", "все категории техники, но один регион"),
        ("C", "все категории и вся Россия сразу"),
    ],
    "Вариант A. Начать с направлений, где уже есть техника и знакомые мастера, проверить процессы, затем расширять каталог и географию.",
    why="Ограниченный запуск позволяет получить реальные объявления и обратную связь без имитации большой пустой площадки.",
)

doc.add_page_break()
add_heading(doc, "Референсы", 1)
add_para(doc, "В завершение встречи попросить показать не «весь сайт как Юла», а несколько конкретных экранов Юлы и сайта «Монитор»:")
add_numbered(doc, "какое поведение нужно повторить;")
add_numbered(doc, "что нравится визуально;")
add_numbered(doc, "что в примере неудобно;")
add_numbered(doc, "что должно работать иначе в нашем проекте.")
add_callout(doc, "Важно:", "внешний вид и расположение элементов проектирует разработчик. Заказчик подтверждает крупные сценарии и ожидаемый результат.", fill=LIGHT_GOLD, label_color=GOLD)

add_summary_table(doc)

add_heading(doc, "Следующий шаг", 1)
add_para(doc, "Аналитик оформляет выбранные варианты в схему ролей и трёх основных процессов: продажа техники, предложение услуги, запрос клиента. Заказчик проверяет схему и подтверждает только бизнес-решения. После этого создаются прототипы экранов и технический план реализации.")

doc.core_properties.title = "Вопросы для обсуждения концепции проекта"
doc.core_properties.subject = "Discovery brief для встречи с заказчиком"
doc.core_properties.author = "Проектная команда"
doc.core_properties.keywords = "требования, доска объявлений, техника, услуги, мастера"
doc.save(OUT_PATH)
print(OUT_PATH)
